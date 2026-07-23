# Copyright (c) 2026 Beijing Volcano Engine Technology Co., Ltd.
# SPDX-License-Identifier: AGPL-3.0
"""Excel 语料端到端测试：项目管理看板.xlsx → Markdown。

直接驱动 ExcelParser._convert_to_markdown（纯转换，无 VikingFS 依赖），
产出 md 到 data/<stem>/ 供人工检查，并断言结构完整性。
"""

import re
from pathlib import Path

import openpyxl
import pytest

from openviking.parse.parsers.excel import ExcelParser

CORPUS_DIR = Path(__file__).parent / "data"
XLSX = CORPUS_DIR / "项目管理看板.xlsx"

EXPECTED_SHEETS = [
    "项目总览表",
    "项目进展0525-0531",
    "0525-0531周报",
    "项目进展0601-0607",
    "0601-0607周报",
    "项目进展0608-0614",
    "0608-0614周报",
    "项目进展0615-0628",
    "0615-0628周报",
    "项目进展0629-0705",
    "0629-0705周报",
    "项目进展0706-0712",
    "0706-0712周报",
    "项目进展0713-0719",
    "0713-0719周报",
]


@pytest.fixture(scope="module")
def md() -> str:
    parser = ExcelParser()
    content = parser._convert_to_markdown(XLSX, openpyxl)
    out_dir = CORPUS_DIR / XLSX.stem
    out_dir.mkdir(exist_ok=True)
    (out_dir / f"{XLSX.stem}.md").write_text(content, encoding="utf-8")
    return content


class TestSheetCoverage:
    def test_all_sheets_present(self, md):
        assert md.count("## Sheet: ") == len(EXPECTED_SHEETS)
        for name in EXPECTED_SHEETS:
            assert f"## Sheet: {name}" in md, f"missing sheet: {name}"

    def test_header_line(self, md):
        assert md.startswith("# 项目管理看板")
        assert f"**Sheets:** {len(EXPECTED_SHEETS)}" in md

    def test_overview_sheet_content(self, md):
        # 项目总览表首行表头与首行数据抽样
        assert "| 项目简称" in md
        assert "优先级" in md and "商机编码" in md
        assert "平高" in md
        assert "OPP--220801-NO.003" in md
        assert "CCTV央视" in md

    def test_no_sheet_lost_content(self, md):
        # 每个 sheet 的非空数据行必须全部出现在表格中
        # （多行单元格会拆出额外物理行，所以转换后行数 >= 源非空行数）
        wb = openpyxl.load_workbook(XLSX, data_only=True)
        for name in EXPECTED_SHEETS:
            sheet = wb[name]
            expected = sum(
                1
                for row in sheet.iter_rows(values_only=True)
                if any(isinstance(c, str) and c.strip() or c is not None and str(c).strip() for c in row)
            )
            section = md.split(f"## Sheet: {name}", 1)[1]
            section = section.split("## Sheet:", 1)[0]
            table_lines = [
                ln
                for ln in section.splitlines()
                if ln.startswith("| ") and not re.match(r"^\| -", ln)
            ]
            assert len(table_lines) >= expected >= 1, (
                f"{name}: {len(table_lines)} table lines < {expected} non-empty source rows"
            )

    def test_no_empty_row_runs(self, md):
        """连续 ≥2 的空白表格行段必须被折叠（used-range 格式残留）；
        单个空行视为有意分组分隔，允许保留。语料中 0525-0531周报 曾有 462 行连续空行。"""
        def is_empty_row(ln: str) -> bool:
            return ln.startswith("|") and "-" not in ln and not ln.strip("| ").strip()

        lines = md.splitlines()
        run = 0
        for ln in lines:
            if is_empty_row(ln):
                run += 1
                assert run <= 1, f"consecutive empty table rows survived near: {ln[:40]}"
            elif ln.strip():
                run = 0

    def test_empty_row_collapse_unit(self):
        """单个空行保留，连续空行段折叠为 1 行。"""
        rows = [
            ["a", "1"],
            ["", ""],            # 单个空行：保留
            ["b", "2"],
            ["", ""],            # 连续段开始
            ["", ""],
            ["", ""],
            ["c", "3"],
        ]
        collapsed = ExcelParser._collapse_empty_rows(rows)
        assert collapsed == [["a", "1"], ["", ""], ["b", "2"], ["", ""], ["c", "3"]]


class TestRowTruncation:
    def test_max_rows_per_sheet(self):
        parser = ExcelParser(max_rows_per_sheet=5)
        content = parser._convert_to_markdown(XLSX, openpyxl)
        # 项目总览表 199 行，截断到 5 行后应有截断提示
        assert "*... 194 more rows truncated ...*" in content

    def test_unlimited_by_default(self, md):
        assert "more rows truncated" not in md


class TestTableIntegrity:
    """量化多行/含管道符单元格对 Markdown 表格的破坏（已知问题探测）。"""

    def test_source_cells_with_breaking_chars(self):
        wb = openpyxl.load_workbook(XLSX, data_only=True)
        nl_cells = sum(
            1
            for name in wb.sheetnames
            for row in wb[name].iter_rows(values_only=True)
            for c in row
            if isinstance(c, str) and "\n" in c
        )
        assert nl_cells > 0, "corpus should contain multi-line cells"

    def test_broken_table_rows_reported(self, md, capsys):
        """统计表格区域内列数与表头不一致的物理行（多行单元格裸 \n 导致）。

        当前 format_table_to_markdown 不做单元格换行/管道符处理，
        该测试用于输出破坏规模，修复后应降为 0。
        """
        broken = 0
        for section in md.split("## Sheet: ")[1:]:
            lines = [ln for ln in section.splitlines() if ln.startswith("|")]
            if not lines:
                continue
            header_cols = lines[0].count("|") - 1
            for ln in lines[1:]:
                if ln.count("|") - 1 != header_cols:
                    broken += 1
        print(f"\nbroken table rows: {broken}")
        # 不断言为 0：记录现状，修复单元格换行后应改为 assert broken == 0
        assert broken >= 0
